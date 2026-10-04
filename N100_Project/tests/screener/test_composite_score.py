"""
Unit tests for Winsorised Composite Scoring and Excel Export (Day 17).
"""

from pathlib import Path
import openpyxl
import pandas as pd
import pytest

from src.screener.engine import (
    load_screener_universe,
    compute_winsorised_composite_score,
    export_screener_excel,
    run_all_presets,
    DEFAULT_EXCEL_PATH,
)


def test_winsorised_composite_score_bounds():
    universe = load_screener_universe()
    scored = compute_winsorised_composite_score(universe)

    assert "composite_quality_score" in scored.columns
    scores = scored["composite_quality_score"].dropna()

    # Scores must be within [0, 100]
    assert (scores >= 0.0).all()
    assert (scores <= 100.0).all()
    assert scores.mean() > 40.0  # Sensible central tendency


def test_screener_excel_generation(tmp_path):
    out_excel = tmp_path / "test_screener_output.xlsx"
    presets = run_all_presets()
    path = export_screener_excel(presets, output_path=out_excel)

    assert path.exists()
    wb = openpyxl.load_workbook(path)

    # 6 Sheets for 6 Presets
    assert len(wb.sheetnames) == 6
    for sname in wb.sheetnames:
        ws = wb[sname]
        # At least header + data rows
        assert ws.max_row >= 2
        # 20 KPI columns
        assert ws.max_column == 20
